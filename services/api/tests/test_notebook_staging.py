"""Staging a notebook assay's workspace in JupyterHub.

``JupyterHub`` starts the user's server and writes files through the contents API;
``plan_files`` lists what ``assay_<id>/`` should hold (the workflow's code, a cohort and
sample-type subset of each measurement input with filtered SDS metadata, whole model
inputs); ``stage_notebook_assay`` writes the entries that are missing.
"""
import io
import sys
from pathlib import Path

import openpyxl
import pytest

project_root = Path(__file__).resolve().parent.parent
if str(project_root) not in sys.path:
    sys.path.append(str(project_root))

from digitaltwins.jupyterhub.hub import HubError, HubTimeout, JupyterHub
from digitaltwins.jupyterhub.staging import Entry, plan_files, stage_notebook_assay, workflow_uuid_by_seek_id

BASE = "http://hub:8000/jupyter"


# ── JupyterHub client ─────────────────────────────────────────────────


class FakeResponse:
    def __init__(self, status_code, body=None):
        self.status_code = status_code
        self._body = body
        self.text = str(body)

    def json(self):
        return self._body


class FakeJupyter:
    """The hub's REST API and user alice's contents API, in memory."""

    def __init__(self, user_exists=True, server=None, polls_until_ready=1):
        self.users = {"alice": {"servers": {"": server} if server else {}}} if user_exists else {}
        self.polls_until_ready = polls_until_ready
        self.dirs = {""}
        self.files = {}
        self.calls = []

    def request(self, method, url, headers=None, json=None, params=None, timeout=None):
        assert headers == {"Authorization": "token tok"}
        assert url.startswith(BASE)
        path = url[len(BASE):]
        self.calls.append((method, path, json))
        if path.startswith("/hub/api/users/alice"):
            return self._hub(method, path[len("/hub/api/users/alice"):])
        prefix = "/user/alice/api/contents/"
        assert path.startswith(prefix), path
        return self._contents(method, path[len(prefix):], json)

    def _hub(self, method, rest):
        if method == "POST" and rest == "":
            self.users["alice"] = {"servers": {}}
            return FakeResponse(201, {"name": "alice"})
        if "alice" not in self.users:
            return FakeResponse(404, {"message": "No such user"})
        servers = self.users["alice"]["servers"]
        if method == "POST" and rest == "/server":
            servers[""] = {"ready": False, "pending": "spawn"}
            return FakeResponse(202)
        if method == "GET" and rest == "":
            server = servers.get("")
            if server and not server["ready"]:
                self.polls_until_ready -= 1
                if self.polls_until_ready <= 0:
                    servers[""] = {"ready": True, "pending": None}
            return FakeResponse(200, {"name": "alice", "servers": dict(servers)})
        raise AssertionError(f"unexpected hub call {method} {rest}")

    def _parent(self, path):
        return path.rsplit("/", 1)[0] if "/" in path else ""

    def _contents(self, method, path, body):
        if method == "GET":
            if path in self.dirs:
                children = [p for p in self.dirs | set(self.files) if p and self._parent(p) == path]
                return FakeResponse(200, {"type": "directory", "content": [
                    {"path": p, "type": "directory" if p in self.dirs else "file"} for p in sorted(children)]})
            if path in self.files:
                return FakeResponse(200, {"type": "file", "path": path})
            return FakeResponse(404)
        if self._parent(path) not in self.dirs:
            return FakeResponse(404, {"message": "parent missing"})
        if method == "PUT" and body["type"] == "directory":
            self.dirs.add(path)
            return FakeResponse(201)
        if method == "PUT":
            assert body["format"] == "base64"
            import base64
            data = base64.b64decode(body["content"])
            chunk = body.get("chunk")
            if chunk is None or chunk == 1:
                self.files[path] = data
            else:
                self.files[path] += data
            return FakeResponse(201)
        if method == "PATCH":
            if body["path"] in self.files:
                return FakeResponse(409)
            self.files[body["path"]] = self.files.pop(path)
            return FakeResponse(200)
        raise AssertionError(f"unexpected contents call {method} {path}")


def hub_for(fake, **kw):
    return JupyterHub(base_url=BASE, token="tok", session=fake, poll_interval=0, **kw)


def test_ensure_server_creates_the_user_and_waits_for_the_spawn():
    fake = FakeJupyter(user_exists=False, polls_until_ready=2)
    hub_for(fake).ensure_server("alice")
    assert ("POST", "/hub/api/users/alice", None) in fake.calls
    assert ("POST", "/hub/api/users/alice/server", None) in fake.calls
    assert fake.users["alice"]["servers"][""]["ready"] is True


def test_ensure_server_leaves_a_running_server_alone():
    fake = FakeJupyter(server={"ready": True, "pending": None})
    hub_for(fake).ensure_server("alice")
    assert [c for c in fake.calls if c[0] != "GET"] == []


def test_ensure_server_times_out_when_the_spawn_never_finishes():
    fake = FakeJupyter(polls_until_ready=10 ** 9)
    with pytest.raises(HubTimeout):
        hub_for(fake, start_timeout=0).ensure_server("alice")


def test_ensure_server_fails_when_the_spawn_disappears():
    fake = FakeJupyter()
    replies = iter([{"": {"ready": False, "pending": "spawn"}}, {}])
    fake._hub = lambda method, rest: FakeResponse(200, {"name": "alice", "servers": next(replies)})
    with pytest.raises(HubError, match="failed to start"):
        hub_for(fake).ensure_server("alice")
    assert [c[0] for c in fake.calls] == ["GET", "GET"]


def test_put_file_uploads_in_parts_under_a_staging_name_then_renames():
    fake = FakeJupyter()
    fake.dirs.add("assay_1")
    hub = hub_for(fake)
    hub.part_size = 3
    hub.put_file("alice", "assay_1/a.bin", iter([b"abc", b"defg"]))
    puts = [(p, b["chunk"]) for m, p, b in fake.calls if m == "PUT"]
    staging = "/user/alice/api/contents/assay_1/a.bin.staging"
    assert puts == [(staging, 1), (staging, 2), (staging, -1)]
    assert fake.files == {"assay_1/a.bin": b"abcdefg"}


def test_put_file_sends_a_small_file_in_one_save():
    fake = FakeJupyter()
    hub_for(fake).put_file("alice", "empty.txt", iter([]))
    [(_, _, body)] = [c for c in fake.calls if c[0] == "PUT"]
    assert "chunk" not in body
    assert fake.files == {"empty.txt": b""}


def test_list_paths_walks_the_tree_and_a_missing_root_is_empty():
    fake = FakeJupyter()
    fake.dirs |= {"assay_1", "assay_1/inputs"}
    fake.files = {"assay_1/nb.ipynb": b"", "assay_1/inputs/x": b"", "other": b""}
    hub = hub_for(fake)
    assert hub.list_paths("alice", "assay_1") == {"assay_1", "assay_1/inputs", "assay_1/nb.ipynb", "assay_1/inputs/x"}
    assert hub.list_paths("alice", "assay_2") == set()


def test_the_token_comes_from_the_environment(monkeypatch):
    monkeypatch.setenv("JUPYTERHUB_API_TOKEN", "tok")
    fake = FakeJupyter(server={"ready": True, "pending": None})
    JupyterHub(base_url=BASE, session=fake).ensure_server("alice")  # FakeJupyter asserts "token tok"


def test_a_missing_token_is_an_error_not_a_default(monkeypatch):
    monkeypatch.delenv("JUPYTERHUB_API_TOKEN", raising=False)
    with pytest.raises(KeyError, match="JUPYTERHUB_API_TOKEN"):
        JupyterHub(base_url=BASE, session=FakeJupyter())


def test_a_failed_call_raises_hub_error():
    fake = FakeJupyter()
    hub = hub_for(fake)
    with pytest.raises(HubError):
        hub.mkdir("alice", "no/parent")


# ── plan_files ────────────────────────────────────────────────────────


def xlsx(rows):
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_of(data):
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=2)]


SUBJECTS = xlsx([["subject id", "age"], ["sub-1", "30y"], ["sub-2", "40y"]])
SAMPLES = xlsx([["sample id", "subject id", "sample type"],
                ["sam-1", "sub-1", "pre"], ["sam-2", "sub-1", "post"], ["sam-1", "sub-2", "pre"]])
MANIFEST = xlsx([["filename", "description"],
                 ["primary/sub-1/sam-1/a.dcm", "a"], ["primary/sub-1/sam-2/b.dcm", "b"],
                 ["primary/sub-2/sam-1/c.dcm", "c"], ["docs/readme.md", "readme"]])


class FakeDownloader:
    def __init__(self, code=True):
        self.buckets = {
            "workflows": {
                "wf-1/primary/workflow_x.cwl": b"cwl",
                **({"wf-1/code/analysis.ipynb": b"{}", "wf-1/code/lib/util.py": b"pass"} if code else {}),
            },
            "measurements": {
                "ds-1/dataset_description.xlsx": b"dd",
                "ds-1/subjects.xlsx": SUBJECTS,
                "ds-1/samples.xlsx": SAMPLES,
                "ds-1/manifest.xlsx": MANIFEST,
                "ds-1/docs/readme.md": b"readme",
                "ds-1/.DS_Store": b"junk",
                "ds-1/primary/sub-1/sam-1/a.dcm": b"a",
                "ds-1/primary/sub-1/sam-2/b.dcm": b"b",
                "ds-1/primary/sub-2/sam-1/c.dcm": b"c",
            },
            "models": {"m-1/primary/model.h5": b"weights"},
        }

    def find_bucket(self, dataset_uuid):
        return next((b for b, objs in self.buckets.items() if any(k.startswith(f"{dataset_uuid}/") for k in objs)), None)

    def list_objects(self, bucket, prefix):
        return sorted(k for k in self.buckets.get(bucket, {}) if k.startswith(prefix))

    def open_object(self, bucket, key):
        body = self.buckets[bucket][key]
        return iter([body]), len(body)


CONFIGS = {
    "inputs": [
        {"name": "#main/raw_dataset", "dataset_uuid": "ds-1", "sample_type": "pre", "category": "measurements"},
        {"name": "weights", "dataset_uuid": "m-1", "sample_type": None, "category": "models"},
    ],
}
SAMPLES_PICKED = [{"subject_id": "sub-1", "sample_id": "sam-1", "dataset_uuid": "ds-1",
                   "sample_type": "pre", "input_name": "#main/raw_dataset"}]


def entries_by_path(entries):
    return {e.path: e for e in entries}


def read(entry):
    return b"".join(entry.open())


def test_plan_holds_the_code_the_sample_subset_whole_models_and_an_outputs_folder():
    entries = entries_by_path(plan_files(CONFIGS, SAMPLES_PICKED, "wf-1", FakeDownloader()))
    assert set(entries) == {
        "analysis.ipynb",
        "lib/util.py",
        "inputs/datasets/raw_dataset/dataset_description.xlsx",
        "inputs/datasets/raw_dataset/subjects.xlsx",
        "inputs/datasets/raw_dataset/samples.xlsx",
        "inputs/datasets/raw_dataset/manifest.xlsx",
        "inputs/datasets/raw_dataset/docs/readme.md",
        "inputs/datasets/raw_dataset/primary/sub-1/sam-1/a.dcm",
        "inputs/models/weights/primary/model.h5",
        "outputs/datasets/",
    }
    assert read(entries["analysis.ipynb"]) == b"{}"
    assert read(entries["inputs/datasets/raw_dataset/primary/sub-1/sam-1/a.dcm"]) == b"a"
    assert read(entries["inputs/models/weights/primary/model.h5"]) == b"weights"
    assert entries["outputs/datasets/"].is_dir


def test_plan_filters_the_sds_metadata_to_the_staged_samples():
    entries = entries_by_path(plan_files(CONFIGS, SAMPLES_PICKED, "wf-1", FakeDownloader()))
    folder = "inputs/datasets/raw_dataset/"
    assert rows_of(read(entries[folder + "subjects.xlsx"])) == [("sub-1", "30y")]
    assert rows_of(read(entries[folder + "samples.xlsx"])) == [("sam-1", "sub-1", "pre")]
    assert rows_of(read(entries[folder + "manifest.xlsx"])) == [
        ("primary/sub-1/sam-1/a.dcm", "a"), ("docs/readme.md", "readme")]


def test_plan_without_workflow_code_is_an_error():
    with pytest.raises(ValueError, match="code"):
        plan_files(CONFIGS, SAMPLES_PICKED, "wf-1", FakeDownloader(code=False))


def test_plan_with_an_input_missing_from_the_object_store_is_an_error():
    configs = {"inputs": [{"name": "x", "dataset_uuid": "gone", "sample_type": "pre", "category": "measurements"}]}
    with pytest.raises(ValueError, match="gone"):
        plan_files(configs, [], "wf-1", FakeDownloader())


# ── stage_notebook_assay ──────────────────────────────────────────────


class FakeHub:
    def __init__(self, existing=()):
        self.existing = set(existing)
        self.log = []

    def ensure_server(self, username):
        self.log.append(("ensure_server", username))

    def list_paths(self, username, root):
        return {p for p in self.existing if p == root or p.startswith(root + "/")}

    def mkdir(self, username, path):
        self.log.append(("mkdir", path))

    def put_file(self, username, path, chunks):
        self.log.append(("put", path, b"".join(chunks)))


def test_stage_writes_missing_entries_after_their_folders_and_skips_existing_ones():
    entries = [
        Entry("nb.ipynb", lambda: [b"new"]),
        Entry("inputs/datasets/raw/a.dcm", lambda: [b"a"]),
        Entry("outputs/datasets/"),
    ]
    hub = FakeHub(existing={"assay_7", "assay_7/nb.ipynb"})
    counts = stage_notebook_assay(hub, "alice", 7, entries)
    assert counts == {"staged": 1, "skipped": 1}
    assert hub.log == [
        ("ensure_server", "alice"),
        ("mkdir", "assay_7/inputs"),
        ("mkdir", "assay_7/inputs/datasets"),
        ("mkdir", "assay_7/inputs/datasets/raw"),
        ("put", "assay_7/inputs/datasets/raw/a.dcm", b"a"),
        ("mkdir", "assay_7/outputs"),
        ("mkdir", "assay_7/outputs/datasets"),
    ]


def test_stage_into_a_new_workspace_creates_the_assay_folder():
    hub = FakeHub()
    stage_notebook_assay(hub, "alice", 7, [Entry("nb.ipynb", lambda: [b"x"])])
    assert hub.log[1:] == [("mkdir", "assay_7"), ("put", "assay_7/nb.ipynb", b"x")]


# ── workflow_uuid_by_seek_id ──────────────────────────────────────────


class FakeConn:
    def __init__(self, rows):
        self.rows = rows
        self.executed = []

    def cursor(self):
        conn = self

        class Cursor:
            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def execute(self, sql, params):
                conn.executed.append((sql, params))

            def fetchall(self):
                return conn.rows

        return Cursor()

    def commit(self):
        pass


def test_workflow_uuid_is_looked_up_among_workflow_datasets_by_seek_id():
    conn = FakeConn([("1a6c77d6-0000-0000-0000-000000000000",)])
    assert workflow_uuid_by_seek_id(conn, 97) == "1a6c77d6-0000-0000-0000-000000000000"
    [(sql, params)] = conn.executed
    assert "workflow_type IS NOT NULL" in sql
    assert params == ("97",)


@pytest.mark.parametrize("seek_id", [97, None])
def test_an_unregistered_or_missing_workflow_is_a_lookup_error(seek_id):
    with pytest.raises(LookupError):
        workflow_uuid_by_seek_id(FakeConn([]), seek_id)
