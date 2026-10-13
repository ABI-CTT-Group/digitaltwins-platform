"""SDS metadata of a script assay's output dataset: prefixed subject/sample ids, and a
manifest that names the files the run actually produced."""
import os
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient

from app.main import app
from app.routers import assays
from app.routers.auth import validate_credentials
from app.routers.dependencies import get_querier, get_uploader

MANIFEST_HEADER = ["filename", "timestamp", "description", "file type", "Additional Metadata"]


def _column(xlsx: Path, header: str) -> list:
    rows = list(openpyxl.load_workbook(xlsx).active.iter_rows(values_only=True))
    i = [str(h).strip().lower() if h else "" for h in rows[0]].index(header)
    return [r[i] for r in rows[1:] if r[i] not in (None, "")]


def _write_manifest(dataset: Path, rows: list) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(MANIFEST_HEADER)
    for r in rows:
        ws.append(r)
    wb.save(dataset / "manifest.xlsx")


def _add_file(dataset: Path, rel: str) -> None:
    path = dataset / "primary" / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"x")


def _manifest_rows(dataset: Path) -> list:
    rows = list(openpyxl.load_workbook(dataset / "manifest.xlsx").active.iter_rows(values_only=True))
    return [r[:3] for r in rows[1:] if r[0]]


# ── _create_sds_output: subject and sample ids ────────────────────────

CONFIGS = {
    "assay_id": 1,
    "outputs": [
        {"name": "nrrd dataset", "dataset_name": "converted_dataset", "sample_name": "nrrd"},
        {"name": "nifti dataset", "dataset_name": "converted_dataset", "sample_name": "nifti"},
    ],
}
SAMPLES = [{"subject_id": "sub-1", "sample_id": "sam-1"}, {"subject_id": "sub-2", "sample_id": "sam-1"}]


def _create(tmp_path: Path) -> Path:
    with patch.object(assays, "MinioUploader", MagicMock()):
        assays._create_sds_output(CONFIGS, SAMPLES, str(tmp_path))
    return tmp_path


def test_output_subjects_use_sub_prefix(tmp_path):
    out = _create(tmp_path)
    assert sorted(_column(out / "subjects.xlsx", "subject id")) == ["sub-1", "sub-2"]


def test_output_samples_use_sub_and_sam_prefixes(tmp_path):
    out = _create(tmp_path)
    assert _column(out / "samples.xlsx", "subject id") == ["sub-1", "sub-1", "sub-2", "sub-2"]
    assert _column(out / "samples.xlsx", "sample id") == ["sam-1", "sam-2", "sam-1", "sam-2"]


# ── _reconcile_manifest: rows name the real files ─────────────────────

def test_reconcile_renames_a_guessed_row_to_the_real_file(tmp_path):
    _add_file(tmp_path, "sub-2/sam-1/ax_t1.nrrd")
    _write_manifest(tmp_path, [["sub-2/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"]])

    assays._reconcile_manifest(str(tmp_path))

    assert _manifest_rows(tmp_path) == [("sub-2/sam-1/ax_t1.nrrd", "t0", "Generated output nrrd")]


def test_reconcile_keeps_a_row_that_already_matches(tmp_path):
    _add_file(tmp_path, "sub-1/sam-1/image.nrrd")
    _write_manifest(tmp_path, [["sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"]])

    assays._reconcile_manifest(str(tmp_path))

    assert _manifest_rows(tmp_path) == [("sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd")]


def test_reconcile_adds_a_row_for_an_unlisted_file(tmp_path):
    _add_file(tmp_path, "sub-1/sam-1/image.nrrd")
    _add_file(tmp_path, "sub-1/sam-1/extra.json")
    _write_manifest(tmp_path, [["sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"]])

    assays._reconcile_manifest(str(tmp_path))

    assert _manifest_rows(tmp_path) == [
        ("sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd"),
        ("sub-1/sam-1/extra.json", "t0", "Generated output nrrd"),
    ]


def test_reconcile_drops_a_row_whose_folder_is_empty(tmp_path):
    _add_file(tmp_path, "sub-1/sam-1/image.nrrd")
    (tmp_path / "primary" / "sub-2" / "sam-1").mkdir(parents=True)
    _write_manifest(tmp_path, [
        ["sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"],
        ["sub-2/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"],
    ])

    assays._reconcile_manifest(str(tmp_path))

    assert _manifest_rows(tmp_path) == [("sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd")]


def test_reconcile_renames_rows_after_a_dropped_row(tmp_path):
    (tmp_path / "primary" / "sub-1" / "sam-1").mkdir(parents=True)
    _add_file(tmp_path, "sub-2/sam-1/ax_t1.nrrd")
    _add_file(tmp_path, "sub-3/sam-1/ax_t2.nrrd")
    _write_manifest(tmp_path, [
        ["sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"],
        ["sub-2/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"],
        ["sub-3/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"],
    ])

    assays._reconcile_manifest(str(tmp_path))

    assert _manifest_rows(tmp_path) == [
        ("sub-2/sam-1/ax_t1.nrrd", "t0", "Generated output nrrd"),
        ("sub-3/sam-1/ax_t2.nrrd", "t0", "Generated output nrrd"),
    ]


def test_reconcile_leaves_a_dataset_without_manifest_alone(tmp_path):
    _add_file(tmp_path, "sub-1/sam-1/image.nrrd")

    assays._reconcile_manifest(str(tmp_path))

    assert not (tmp_path / "manifest.xlsx").exists()


# ── upload_workspace_datasets publishes the reconciled manifest ───────

def test_submit_publishes_the_manifest_of_the_real_files():
    published = {}

    def upload_dataset(dataset_path, category):
        published["rows"] = _manifest_rows(Path(dataset_path))
        return "uuid-1"

    uploader = MagicMock()
    uploader.upload_dataset.side_effect = upload_dataset
    querier = MagicMock()
    querier.get_assay.return_value = {"configs": {"outputs": [
        {"dataset_name": "converted_dataset", "category": "measurements"}]}}

    def download_folder(bucket, prefix, save_dir):
        dataset = Path(save_dir) / "converted_dataset"
        _add_file(dataset, "sub-1/sam-1/ax_t1.nrrd")
        _write_manifest(dataset, [["sub-1/sam-1/image.nrrd", "t0", "Generated output nrrd", "image"]])
        return 1

    saved = dict(app.dependency_overrides)
    app.dependency_overrides[validate_credentials] = lambda: {"username": "u", "token": "t"}
    app.dependency_overrides[get_uploader] = lambda: uploader
    app.dependency_overrides[get_querier] = lambda: querier
    try:
        with patch("digitaltwins.minio.downloader.Downloader") as Downloader:
            Downloader.return_value.get_latest_timestamp_folder.return_value = "20261013_000000"
            Downloader.return_value.download_folder.side_effect = download_folder
            response = TestClient(app).post("/assays/1/workspace/dataset/upload")
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(saved)

    assert response.status_code == 200, response.text
    assert published["rows"] == [("sub-1/sam-1/ax_t1.nrrd", "t0", "Generated output nrrd")]
