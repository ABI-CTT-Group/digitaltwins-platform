"""Stage a notebook assay's workspace, ``assay_<id>/`` in the user's JupyterHub server.

The folder gets the workflow's ``code/``, a cohort and sample-type subset of each measurement
input under ``inputs/datasets/<input>/`` (with its SDS metadata filtered to match), each model
input whole under ``inputs/models/<input>/``, and an empty ``outputs/datasets/``, which the
workspace download and upload endpoints read. Only missing files are written, so the user's
edits survive a re-launch and an interrupted launch resumes.
See docs/decisions/2026-10-16-stage-notebook-assay-workspace.md.
"""
import io
from dataclasses import dataclass
from typing import Callable, Iterable, List, Optional, Set, Tuple

import openpyxl

from .. import models, workflows

# SDS metadata files that list subjects, samples or files, filtered to the staged samples.
_FILTERED = ("subjects.xlsx", "samples.xlsx", "manifest.xlsx")


@dataclass(frozen=True)
class Entry:
    """A path under ``assay_<id>/``: a folder when it ends with ``/``, else a file read by ``open``."""
    path: str
    open: Optional[Callable[[], Iterable[bytes]]] = None

    @property
    def is_dir(self) -> bool:
        return self.path.endswith("/")


def workflow_uuid_by_seek_id(conn, seek_id) -> str:
    """The platform workflow dataset registered as SEEK workflow ``seek_id``."""
    if not seek_id:
        raise LookupError("The assay has no linked workflow.")
    with conn.cursor() as cur:
        cur.execute("SELECT dataset_uuid FROM dataset WHERE seek_id = %s AND workflow_type IS NOT NULL",
                    (str(seek_id),))
        rows = cur.fetchall()
    conn.commit()
    if not rows:
        raise LookupError(f"Workflow {seek_id} is not registered in the platform.")
    return str(rows[0][0])


def _input_folder(name: str) -> str:
    """The input's name without a CWL ``#main/`` prefix."""
    return str(name or "input").rsplit("/", 1)[-1].lstrip("#") or "input"


def _hidden(rel: str) -> bool:
    return any(part.startswith(".") for part in rel.split("/"))


def _object(downloader, bucket: str, key: str, path: str) -> Entry:
    return Entry(path, lambda: downloader.open_object(bucket, key)[0])


def plan_files(configs: dict, samples: List[dict], workflow_uuid: str, downloader) -> List[Entry]:
    """What ``assay_<id>/`` should hold, as paths relative to it.

    ``samples`` are the assay's selected samples (``_discover_samples``). Hidden files are left
    out, as the Jupyter contents API refuses them.
    """
    code_prefix = f"{workflow_uuid}/code/"
    code = [k for k in downloader.list_objects(workflows.CATEGORY, code_prefix) if not k.endswith("/")]
    if not code:
        raise ValueError(f"Workflow dataset {workflow_uuid} has no code/ to stage.")
    entries = [_object(downloader, workflows.CATEGORY, k, k[len(code_prefix):])
               for k in code if not _hidden(k[len(code_prefix):])]

    for inp in configs.get("inputs", []):
        dataset_uuid = inp.get("dataset_uuid")
        if not dataset_uuid:
            continue
        is_model = inp.get("category") == models.CATEGORY
        folder = f"inputs/{'models' if is_model else 'datasets'}/{_input_folder(inp.get('name'))}/"
        bucket = downloader.find_bucket(dataset_uuid)
        if bucket is None:
            raise ValueError(f"Input dataset {dataset_uuid} is not in the object store.")
        rels = [k[len(dataset_uuid) + 1:] for k in downloader.list_objects(bucket, f"{dataset_uuid}/")
                if not k.endswith("/")]
        rels = [r for r in rels if not _hidden(r)]
        if not is_model:
            pairs = {(s["subject_id"], s["sample_id"]) for s in samples
                     if s.get("input_name", "input") == inp.get("name", "input")}
            prefixes = tuple(f"primary/{sub}/{sam}/" for sub, sam in pairs)
            rels = [r for r in rels if not r.startswith("primary/") or r.startswith(prefixes)]
        for rel in rels:
            key = f"{dataset_uuid}/{rel}"
            if not is_model and rel in _FILTERED:
                entries.append(Entry(folder + rel, _filtered(downloader, bucket, key, rel, pairs, set(rels))))
            else:
                entries.append(_object(downloader, bucket, key, folder + rel))

    entries.append(Entry("outputs/datasets/"))
    return entries


def _filtered(downloader, bucket: str, key: str, name: str,
              pairs: Set[Tuple[str, str]], staged: Set[str]) -> Callable[[], Iterable[bytes]]:
    def read() -> Iterable[bytes]:
        chunks, _ = downloader.open_object(bucket, key)
        return [filter_metadata(name, b"".join(chunks), pairs, staged)]
    return read


def filter_metadata(name: str, data: bytes, pairs: Set[Tuple[str, str]], staged: Set[str]) -> bytes:
    """Drop the rows of ``subjects``/``samples``/``manifest`` xlsx that name unstaged subjects, samples or files.

    Manifest rows outside ``primary/`` are kept; rows inside it are kept when their file is staged.
    """
    subjects = {sub for sub, _ in pairs}
    if name == "subjects.xlsx":
        def keep(row):
            return row.get("subject id") in subjects
    elif name == "samples.xlsx":
        def keep(row):
            return (row.get("subject id"), row.get("sample id")) in pairs
    else:
        def keep(row):
            return bool(row.get("filename")) and (
                not row["filename"].startswith("primary/") or row["filename"] in staged)

    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    for r in range(ws.max_row, 1, -1):
        row = {h: str(c.value).strip() for h, c in zip(header, ws[r]) if c.value is not None}
        if not keep(row):
            ws.delete_rows(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def stage_notebook_assay(hub, username: str, assay_id, entries: List[Entry]) -> dict:
    """Write the entries missing from ``assay_<id>/``; return ``{staged, skipped}`` file counts."""
    root = f"assay_{assay_id}"
    hub.ensure_server(username)
    existing = hub.list_paths(username, root)
    staged = skipped = 0
    for entry in entries:
        path = f"{root}/{entry.path}".rstrip("/")
        if path in existing:
            skipped += not entry.is_dir
            continue
        _makedirs(hub, username, path if entry.is_dir else path.rsplit("/", 1)[0], existing)
        if entry.is_dir:
            continue
        hub.put_file(username, path, entry.open())
        existing.add(path)
        staged += 1
    return {"staged": staged, "skipped": skipped}


def _makedirs(hub, username: str, path: str, existing: Set[str]) -> None:
    parts = path.split("/")
    for i in range(1, len(parts) + 1):
        folder = "/".join(parts[:i])
        if folder not in existing:
            hub.mkdir(username, folder)
            existing.add(folder)
